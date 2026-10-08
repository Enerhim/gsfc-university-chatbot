import json,sys,subprocess,hashlib,io,shutil
from pathlib import Path
from pypdf import PdfReader
from .store import DATA,ingest,record

def process(path):
    obj=json.loads(Path(path).read_text());s=obj.get('structuredContent') or {};mid=s.get('message_id','');name=s.get('filename','');url='https://mail.google.com/mail/u/0/#all/'+mid
    key=hashlib.sha256((mid+name).encode()).hexdigest()[:20];status='failed';text='';detail=''
    output=DATA/'attachments';output.mkdir(exist_ok=True,mode=0o700)
    try:
        uri=s.get('file_uri',{});download=uri.get('download_url','') if isinstance(uri,dict) else uri
        suffix=Path(name).suffix.lower();dest=output/(key+suffix)
        if download.startswith('https://'):
            r=subprocess.run(['curl','-fLsS','--max-time','60','--max-filesize','31457280',download,'-o',str(dest)],capture_output=True)
            if r.returncode:raise RuntimeError('Attachment download failed')
            if suffix=='.pdf':
                pdf=PdfReader(dest);text='\n\n'.join('[Page '+str(i+1)+']\n'+(p.extract_text() or '') for i,p in enumerate(pdf.pages))
            elif suffix=='.docx':
                from docx import Document
                doc=Document(dest);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
            elif suffix in ('.txt','.csv','.ics'):text=dest.read_text(errors='replace')
            elif suffix in ('.xlsx','.xls'):
                try:
                    import openpyxl
                    wb=openpyxl.load_workbook(dest,read_only=True,data_only=True);text='\n'.join(s.title+'\n'+'\n'.join(' | '.join(str(c or '') for c in row) for row in s.values) for s in wb)
                except Exception:pass
            elif suffix in ('.png','.jpg','.jpeg','.tif','.tiff'):
                from .ocr import image_text
                text=image_text(dest)
        if not text.strip() and not s.get('content_truncated'):
            text='\n'.join(c.get('text','') for c in s.get('content',[]) if c.get('type')=='text')
        if len(text.strip())>=40:
            ingest(url+'?attachment='+key,name,text,'attachment','private',metadata={'message_id':mid,'filename':name,'local_file':str(dest) if download else ''});status='indexed'
        else:status='needs_ocr_or_extraction'
    except Exception as e:detail=str(e)[:200]
    with (DATA/'attachment-status.jsonl').open('a') as f:f.write(json.dumps({'message_id':mid,'filename':name,'key':key,'status':status,'detail':detail})+'\n')
    Path(path).unlink(missing_ok=True)
    return status
if __name__=='__main__':print(process(sys.argv[1]))
